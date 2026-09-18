from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import threading
from functools import lru_cache
import uuid
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from flask import Flask, Response, jsonify, redirect, render_template, request, url_for
from markupsafe import Markup


ROOT = Path(__file__).resolve().parent
DEFAULT_DB = ROOT / "price_history.sqlite3"
DEFAULT_CONFIG = ROOT / "config.json"
HOSTED_READ_ONLY = bool(os.environ.get("VERCEL"))
DISPLAY_TIMEZONE = ZoneInfo("Asia/Kolkata")
SUPPORTED_BRANDS = [
    "Apple", "Asus", "Google", "Honor", "Infinix", "iQOO", "Mi", "Motorola",
    "Nokia", "Nothing", "OnePlus", "OPPO", "POCO", "realme", "Samsung", "vivo", "Xiaomi",
]
SCRAPE_BRANDS = {
    "all", "apple", "samsung", "oppo", "realme", "oneplus", "mi", "vivo", "iqoo", "motorola", "cashify",
}

app = Flask(__name__)
job_lock = threading.Lock()
job_state: dict[str, Any] = {
    "id": None,
    "running": False,
    "started_at": None,
    "finished_at": None,
    "returncode": None,
    "message": "Ready",
    "output": "",
    "progress": {
        "done": 0,
        "rows": 0,
        "errors": 0,
        "current": "",
        "scope": "",
    },
}


def db_path() -> Path:
    configured = app.config.get("DB_PATH") or os.environ.get("DB_PATH")
    if not configured:
        return DEFAULT_DB
    path = Path(configured)
    if not path.exists() and DEFAULT_DB.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(DEFAULT_DB, path)
    return path


def connect() -> sqlite3.Connection:
    conn = sqlite3.connect(db_path())
    conn.row_factory = sqlite3.Row
    ensure_schema(conn)
    return conn


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS price_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date TEXT NOT NULL,
            brand TEXT NOT NULL,
            model TEXT NOT NULL,
            part TEXT NOT NULL,
            price TEXT,
            price_value REAL,
            currency TEXT,
            url TEXT NOT NULL,
            status TEXT NOT NULL,
            error TEXT
        )
        """
    )
    conn.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_price_history_lookup
        ON price_history (brand, model, part, date)
        """
    )
    conn.commit()


def money(value: float | None, currency: str | None) -> str:
    if value is None:
        return "-"
    prefix = "Rs " if currency == "INR" else f"{currency} " if currency else ""
    return f"{prefix}{value:,.0f}"


@app.template_filter("money")
def money_filter(value: float | None, currency: str | None = None) -> str:
    return money(value, currency)


ICONS: dict[str, str] = {
    "refresh": '<path d="M3 12a9 9 0 0 1 15.3-6.4M21 12a9 9 0 0 1-15.3 6.4"/><path d="M3 3v5h5M21 21v-5h-5"/>',
    "download": '<path d="M12 3v12m0 0-4-4m4 4 4-4"/><path d="M4 19h16"/>',
    "dot": '<circle cx="12" cy="12" r="5"/>',
    "play": '<path d="M6 4l14 8-14 8V4z"/>',
    "search": '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/>',
    "search-sm": '<circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/>',
    "target": '<circle cx="12" cy="12" r="8"/><circle cx="12" cy="12" r="4"/><circle cx="12" cy="12" r="0.5"/>',
    "box": '<path d="M21 8 12 3 3 8v8l9 5 9-5V8Z"/><path d="M3 8l9 5 9-5M12 13v8"/>',
    "layers": '<path d="m12 2 9 5-9 5-9-5 9-5Z"/><path d="m3 12 9 5 9-5M3 17l9 5 9-5"/>',
    "clock": '<circle cx="12" cy="12" r="9"/><path d="M12 7v5l3 3"/>',
    "tag": '<path d="M20 12 12.5 19.5a2 2 0 0 1-2.8 0l-6.2-6.2a2 2 0 0 1 0-2.8L11 3h9v9Z"/><circle cx="15" cy="8" r="1.5"/>',
    "chevron": '<path d="m9 6 6 6-6 6"/>',
    "empty": '<path d="M4 7h16l-1.5 12.5a2 2 0 0 1-2 1.8H7.5a2 2 0 0 1-2-1.8L4 7Z"/><path d="M9 7V5a3 3 0 0 1 6 0v2"/>',
    "check": '<path d="m5 12 5 5 9-10"/>',
    "alert": '<path d="M12 3 2 21h20L12 3Z"/><path d="M12 10v4M12 17h.01"/>',
    "asc": '<path d="m6 15 6-6 6 6"/>',
    "desc": '<path d="m6 9 6 6 6-6"/>',
    "sort-none": '<path d="m8 9 4-4 4 4M8 15l4 4 4-4" opacity="0.4"/>',
}


@app.template_global("icon")
def icon(name: str) -> str:
    body = ICONS.get(name, ICONS["dot"])
    svg = f'<svg class="icon icon-{name}" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true">{body}</svg>'
    return Markup(svg)


@app.template_filter("localdate")
def localdate_filter(value: str | None) -> str:
    if not value:
        return "-"
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        local_time = parsed.astimezone(DISPLAY_TIMEZONE)
        return local_time.strftime("%d %b %Y, %I:%M %p")
    except ValueError:
        return value


PER_PAGE_OPTIONS = [25, 50, 100, 250]
SORT_KEYS: dict[str, Any] = {
    "brand": lambda r: (r["brand"] or "").lower(),
    "model": lambda r: r["model_key"],
    "part": lambda r: (r["part"] or "").lower(),
    "price": lambda r: (r["price_value"] if r["price_value"] is not None else -1),
    "date": lambda r: r["date"] or "",
}


COMPETITOR_BRAND_ALIASES = {"xiaomi": "mi", "poco": "mi"}
MODEL_PREFIX_RE = re.compile(r"^(?:xt\d{4}(?:-\d+)?\s+)?(?:(?:apple|samsung|xiaomi|oppo|realme|oneplus|vivo|iqoo|motorola|moto)\s+)*")
IPHONE_SE_RE = re.compile(r"iphone se\s*\(?(\d)(?:st|nd|rd|th) generation\)?")
IPHONE_SE_YEARS = {"2016": "1", "2020": "2", "2022": "3"}
LOOSE_SUFFIX_RE = re.compile(r"\s*\(?5g\)?$")


@lru_cache(maxsize=None)
def clean_text(value: str | None) -> str:
    """Lowercase, replace non-breaking spaces (Apple's site uses them inside
    model names) and collapse runs of whitespace."""
    return re.sub(r"\s+", " ", (value or "").replace("\u00a0", " ")).strip().lower()


@lru_cache(maxsize=None)
def clean_label(value: str | None) -> str:
    """Display text with non-breaking spaces and doubled whitespace removed."""
    return re.sub(r"\s+", " ", (value or "").replace("\u00a0", " ")).strip()


def brand_key(brand: str | None) -> str:
    key = clean_text(brand)
    return COMPETITOR_BRAND_ALIASES.get(key, key)


@lru_cache(maxsize=None)
def model_key(model: str | None, loose: bool = False) -> str:
    """Canonical model name shared by every source: strips the brand prefix
    Cashify adds ("Apple iPhone 14" -> "iphone 14"), Motorola's XT codes,
    and folds iPhone SE year/generation spellings together. With loose=True
    a trailing "5G" is dropped too, for sources that omit it."""
    key = MODEL_PREFIX_RE.sub("", clean_text(model))
    key = IPHONE_SE_RE.sub(lambda m: f"iphone se gen {m.group(1)}", key)
    key = re.sub(
        r"iphone se (2016|2020|2022)$",
        lambda m: f"iphone se gen {IPHONE_SE_YEARS[m.group(1)]}",
        key,
    )
    if key == "iphone se":
        key = "iphone se gen 1"
    # Samsung writes "Galaxy Z Fold7"; Cashify writes "Galaxy Z Fold 7".
    key = re.sub(r"\b(fold|flip|note)\s+(?=\d)", r"\1", key)
    if loose:
        key = LOOSE_SUFFIX_RE.sub("", key)
    return key


@lru_cache(maxsize=None)
def _search_text(brand: str, model: str, part: str, price: str, extra: tuple[str, ...]) -> str:
    return clean_text(" ".join([brand, model, part, price, model_key(model), *extra]))


def search_text(row: Any, *extra: str) -> str:
    """Text a search needle is matched against: the raw fields plus the
    canonical model name, so "iPhone SE (3rd generation)" also finds
    Cashify's "Apple iPhone SE 2022"."""
    return _search_text(
        str(row["brand"] or ""), str(row["model"] or ""), str(row["part"] or ""), str(row["price"] or ""),
        tuple(str(item) for item in extra),
    )


SEARCH_BRAND_WORDS = {
    "apple", "samsung", "xiaomi", "oppo", "realme", "oneplus", "vivo", "iqoo", "motorola", "moto",
}


def matches_search(row: Any, search: str, *extra: str) -> bool:
    needle = clean_text(search)
    if needle in search_text(row, *extra):
        return True
    # Fall back to the canonical model name, so "Apple iPhone SE 2022" finds
    # "iPhone SE (3rd generation)". A brand named in the search must still
    # match: "OPPO A12" must not find Samsung's "Galaxy A12".
    key = model_key(search)
    if not key or key not in model_key(row["model"]):
        return False
    prefix = needle[: -len(key)].split() if needle.endswith(key) else []
    named = {"motorola" if word == "moto" else brand_key(word) for word in prefix if word in SEARCH_BRAND_WORDS}
    return not named or brand_key(row["brand"]) in named


def price_source(row: Any) -> str:
    """Name of the site a price was scraped from: the brand's own support
    site, or Cashify (the third-party repair competitor)."""
    part = row["part"] or ""
    url = row["url"] or ""
    if part.startswith("Cashify - ") or "cashify.in" in url:
        return "Cashify"
    return row["brand"] or "Official"


@lru_cache(maxsize=None)
def part_category(part: str | None) -> str | None:
    """Fold the many vendor-specific spare-part names into a shared category so
    the same repair can be compared across sources (e.g. Apple "Screen damage",
    Mi "LCD Module - Display" and "Cashify - Screen" are all "Screen").
    Returns None for accessories and parts no competitor sells."""
    name = (part or "").lower()
    if name.startswith("cashify - "):
        name = name[len("cashify - "):]
    name = name.replace("-", " ").replace("（", " (").replace("）", ")")

    def has(*words: str) -> bool:
        return any(word in name for word in words)

    if has("cable", "adapter", "charger", "headset", "cleaning", "s pen", "sim tray", "component repair"):
        return None
    if "camera" in name:
        return "Front Camera" if "front" in name else "Back Camera"
    if has("proximity"):
        return "Proximity Sensor"
    if has("aux", "headphone", "earphone jack"):
        return "Aux Jack"
    if has("receiver", "earpiece"):
        return "Receiver"
    if has("speaker"):
        return "Speaker"
    if has("mic"):
        return "Mic"
    if has("charging", "usb port", "sub board", "pcb sb"):
        return "Charging Jack"
    if has("motherboard", "mainboard", "main board", "mother board", "pcb mb", "logic board"):
        return "Motherboard"
    if has("sub screen", "cover screen", "cover display", "outer screen", "outer display", "cli display", "secondary display"):
        return "Cover Screen"
    if has("screen", "display", "lcd", "lcm", "touch panel"):
        return "Screen"
    if has("back cover", "back glass", "battery cover", "backcover", "rear cover", "rear glass"):
        return "Back Cover"
    if has("battery"):
        return "Battery"
    return None


def is_recent_enough(date_value: str | None, newest: str, window_hours: int = 36) -> bool:
    """True when a check happened within `window_hours` of the newest check."""
    try:
        checked = datetime.fromisoformat((date_value or "").replace("Z", "+00:00"))
        latest = datetime.fromisoformat(newest.replace("Z", "+00:00"))
    except ValueError:
        return True
    return (latest - checked) <= timedelta(hours=window_hours)


def attach_competitor_prices(latest: list[dict[str, Any]]) -> None:
    """For every latest price, list the latest prices for the same model and
    part category scraped from a *different* source, so a row can show the
    official and the Cashify price side by side."""
    strict: defaultdict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    loose: defaultdict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in latest:
        row["source"] = price_source(row)
        row["brand_key"] = brand_key(row["brand"])
        row["model_key"] = model_key(row["model"])
        row["category"] = part_category(row["part"])
        row["competitors"] = []
        if row["category"]:
            strict[(row["brand_key"], row["model_key"], row["category"])].append(row)
            loose[(row["brand_key"], model_key(row["model"], loose=True), row["category"])].append(row)

    def others(candidates: list[dict[str, Any]], row: dict[str, Any]) -> list[dict[str, Any]]:
        seen: set[tuple[str, str]] = set()
        found = []
        for other in candidates:
            tag = (other["source"], other["part"])
            if other["source"] == row["source"] or tag in seen:
                continue
            seen.add(tag)
            found.append(other)
        # A source can carry the same repair under two part names (the old
        # config.json shortlist saved Apple's battery as "Battery", discovery
        # saves it as "Battery service"). Keep only the parts that source
        # refreshed in its most recent check, so stale names drop out.
        newest: dict[str, str] = {}
        for other in found:
            newest[other["source"]] = max(newest.get(other["source"], ""), other["date"] or "")
        found = [other for other in found if is_recent_enough(other["date"], newest[other["source"]])]
        return sorted(found, key=lambda r: (r["source"], r["part"]))

    for row in latest:
        if not row["category"]:
            continue
        found = others(strict[(row["brand_key"], row["model_key"], row["category"])], row)
        if not found:
            found = others(loose[(row["brand_key"], model_key(row["model"], loose=True), row["category"])], row)
        row["competitors"] = found


def group_by_model(rows: list[dict[str, Any]]) -> list[tuple[dict[str, Any], bool, int, str, int, str]]:
    """Tag consecutive rows sharing (brand, model) so the template can collapse
    them into one accordion instead of repeating the brand/model on every
    spare-part row. Grouping only merges rows that are already adjacent in the
    current sort order (e.g. sort by Brand or Model) -- it's a display
    annotation, not a re-sort, so other sort orders mostly yield groups of 1
    and render as plain rows.
    """
    keys = [(row["brand_key"], row["model_key"]) for row in rows]
    grouped: list[tuple[dict[str, Any], bool, int, str, int, str]] = []
    group_index = -1
    i = 0
    n = len(rows)
    while i < n:
        j = i
        while j < n and keys[j] == keys[i]:
            j += 1
        group_index += 1
        group_id = f"g{group_index}"
        size = j - i
        sources = len({rows[k]["source"] for k in range(i, j)})
        # Title the accordion with the brand's own model name when we have it;
        # Cashify prefixes the brand ("Apple iPhone 14") and spells some models
        # differently.
        label = next((rows[k]["model"] for k in range(i, j) if rows[k]["source"] != "Cashify"), rows[i]["model"])
        for k in range(i, j):
            grouped.append((rows[k], k == i, size, group_id, sources, label))
        i = j
    return grouped


_PRICE_DATA_LOCK = threading.Lock()
_PRICE_DATA: dict[str, Any] = {"stamp": None}


def load_price_data() -> dict[str, Any]:
    """Everything that depends only on the database contents: all rows, the
    latest price per part with its competitor prices, and the picker lists.

    Building this is the expensive part of a page load (seconds of CPU on a
    small hosted instance), so it is kept until the database file changes.
    Callers must treat the returned lists and rows as read-only."""
    path = db_path()
    try:
        info = path.stat()
        stamp = (str(path), info.st_mtime_ns, info.st_size)
    except OSError:
        stamp = (str(path), 0, 0)
    with _PRICE_DATA_LOCK:
        if _PRICE_DATA.get("stamp") == stamp:
            return _PRICE_DATA
        conn = connect()
        rows = conn.execute(
            """
            SELECT date, brand, model, part, price, price_value, currency, status, error, url
            FROM price_history
            ORDER BY date DESC
            """
        ).fetchall()
        conn.close()

        latest_by_key: dict[tuple[str, str, str], sqlite3.Row] = {}
        for row in reversed(rows):
            if row["status"] == "ok":
                latest_by_key[(row["brand"], row["model"], row["part"])] = row

        latest = [dict(row) for row in latest_by_key.values()]
        attach_competitor_prices(latest)
        latest.sort(key=lambda r: (r["brand_key"], r["model_key"], r["source"] == "Cashify", r["brand"], r["part"]))
        # One name per phone for the export picker; rows are ordered
        # official-first, so setdefault keeps the brand's own spelling.
        model_names: dict[tuple[str, str], str] = {}
        for row in latest:
            model_names.setdefault((row["brand_key"], row["model_key"]), clean_label(row["model"]))

        _PRICE_DATA.update(
            stamp=stamp,
            rows=rows,
            latest=latest,
            brands=sorted(set(SUPPORTED_BRANDS) | {row["brand"] for row in rows}),
            model_options=sorted(set(model_names.values()), key=str.lower),
            suggestions=sorted({f"{row['model']} - {row['part']}" for row in latest})[:500],
        )
        return _PRICE_DATA


def load_dashboard(
    search: str = "",
    selected_brands: list[str] | None = None,
    page: int = 1,
    per_page: int = 100,
    status: str = "",
    sort: str = "date",
    sort_dir: str = "desc",
    paginate: bool = True,
) -> dict[str, Any]:
    selected_brands = selected_brands or []
    selected_brands_lower = {b.lower() for b in selected_brands}
    cached = load_price_data()
    rows = cached["rows"]
    all_latest = cached["latest"]
    latest = list(all_latest)  # sorted and filtered per request; never mutate the cached list
    brands = cached["brands"]
    model_options = cached["model_options"]
    if selected_brands_lower:
        latest = [row for row in latest if row["brand"].lower() in selected_brands_lower]
        rows = [row for row in rows if row["brand"].lower() in selected_brands_lower]
    if status in {"ok", "error"}:
        rows = [row for row in rows if row["status"] == status]
    if search:
        latest = [row for row in latest if matches_search(row, search)]
        rows = [row for row in rows if matches_search(row, search, row["status"])]

    suggestions = cached["suggestions"]

    sort_key = SORT_KEYS.get(sort, SORT_KEYS["date"])
    latest.sort(key=sort_key, reverse=(sort_dir != "asc"))
    if sort not in SORT_KEYS or sort == "date":
        # Keep every model's rows together (official + Cashify were scraped at
        # different times) while ordering models by their most recent check.
        first_seen: dict[tuple[str, str], int] = {}
        for row in latest:
            first_seen.setdefault((row["brand_key"], row["model_key"]), len(first_seen))
        latest.sort(key=lambda r: first_seen[(r["brand_key"], r["model_key"])])

    latest_total = len(latest)
    per_page = per_page if per_page in PER_PAGE_OPTIONS else 100
    if not paginate:
        per_page = max(latest_total, 1)
        page = 1
    total_pages = max(1, (latest_total + per_page - 1) // per_page)
    page = min(max(page, 1), total_pages)
    page_start = (page - 1) * per_page
    page_end = page_start + per_page
    paged_latest = group_by_model(latest[page_start:page_end])

    chart_limit = 24 if search else 12
    chart_candidates = sorted(latest, key=lambda r: r["date"], reverse=True)[:chart_limit]
    chart_keys = {(row["brand"], row["model"], row["part"]) for row in chart_candidates}
    history_by_key: defaultdict[tuple[str, str, str], list[sqlite3.Row]] = defaultdict(list)
    for row in reversed(rows):
        key = (row["brand"], row["model"], row["part"])
        if row["status"] == "ok" and key in chart_keys:
            history_by_key[key].append(row)

    chart_series = []
    for key in chart_keys:
        points = history_by_key.get(key, [])
        chart_series.append(
            {
                "label": " - ".join(key),
                "points": [
                    {"date": row["date"], "value": row["price_value"], "price": money(row["price_value"], row["currency"])}
                    for row in points
                    if row["price_value"] is not None
                ],
            }
        )
    chart_series.sort(key=lambda item: item["label"])

    successful_rows = [row for row in rows if row["status"] == "ok"]
    error_rows = [row for row in rows if row["status"] != "ok"]
    total_value = sum(row["price_value"] or 0 for row in latest)
    priced = [row["price_value"] for row in latest if row["price_value"] is not None]
    average_value = (sum(priced) / len(priced)) if priced else None
    today_local = datetime.now(DISPLAY_TIMEZONE).date()
    updated_today = sum(1 for row in all_latest if _is_local_date(row["date"], today_local))
    chart_note = f"Showing {len(chart_series)} chart series"
    if len(latest) > len(chart_series):
        chart_note += f" from {len(latest)} matching prices. Search a model or spare part for a focused chart."

    window = 2
    page_numbers = sorted(
        {p for p in range(page - window, page + window + 1) if 1 <= p <= total_pages}
        | {1, total_pages}
    )

    return {
        "latest": paged_latest,
        "history": rows[:100],
        "brands": brands,
        "per_page_options": PER_PAGE_OPTIONS,
        "suggestions": suggestions,
        "model_options": model_options,
        "sort": sort,
        "sort_dir": sort_dir,
        "selected_status": status,
        "selected_brands": selected_brands,
        "pagination": {
            "page": page,
            "per_page": per_page,
            "total": latest_total,
            "total_pages": total_pages,
            "start": page_start + 1 if latest_total else 0,
            "end": min(page_end, latest_total),
            "has_prev": page > 1,
            "has_next": page < total_pages,
            "prev_page": max(1, page - 1),
            "next_page": min(total_pages, page + 1),
            "page_numbers": page_numbers,
        },
        "chart_json": json.dumps(chart_series),
        "chart_note": chart_note,
        "stats": {
            "tracked": len(latest),
            "brands": len({row["brand"] for row in latest}),
            "models": len({(row["brand_key"], row["model_key"]) for row in latest}),
            "checks": len(rows),
            "errors": len(error_rows),
            "total_value": total_value,
            "average_value": average_value,
            "updated_today": updated_today,
            "last_success": successful_rows[0]["date"] if successful_rows else None,
        },
        "job": current_job_state(),
        "hosted_read_only": HOSTED_READ_ONLY,
    }


def _is_local_date(value: str | None, target) -> bool:
    if not value:
        return False
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return False
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(DISPLAY_TIMEZONE).date() == target


def parse_filters() -> dict[str, Any]:
    brands = []
    for value in request.args.getlist("brand"):
        cleaned = value.strip()
        if cleaned and cleaned not in brands:
            brands.append(cleaned)
    return {
        "search": request.args.get("q", "").strip(),
        "brands": brands,
        "status": request.args.get("status", "").strip(),
        "page": parse_positive_int(request.args.get("page"), 1),
        "per_page": parse_positive_int(request.args.get("per_page"), 100),
        "sort": request.args.get("sort", "date").strip() or "date",
        "sort_dir": request.args.get("dir", "desc").strip() or "desc",
    }


def build_query(filters: dict[str, Any], **overrides: Any) -> dict[str, Any]:
    """Base query dict (no page) used for sort/export links, which append
    their own page number. Pass page=... in overrides when one is needed,
    e.g. resetting to page 1 after a filter change.
    """
    query: dict[str, Any] = {
        "q": filters["search"],
        "brand": filters["brands"],
        "status": filters["status"],
        "per_page": filters["per_page"],
        "sort": filters["sort"],
        "dir": filters["sort_dir"],
    }
    query.update(overrides)
    return query


def build_active_filters(filters: dict[str, Any]) -> list[dict[str, str]]:
    active = []
    if filters["search"]:
        active.append(
            {
                "label": f'Search "{filters["search"]}"',
                "remove_url": url_for("index", **build_query(filters, q="", page=1)),
            }
        )
    for selected in filters["brands"]:
        remaining = [b for b in filters["brands"] if b != selected]
        active.append(
            {"label": selected, "remove_url": url_for("index", **build_query(filters, brand=remaining, page=1))}
        )
    if filters["status"] in {"ok", "error"}:
        label = "Available" if filters["status"] == "ok" else "Error"
        active.append({"label": label, "remove_url": url_for("index", **build_query(filters, status="", page=1))})
    return active


def warm_price_data() -> None:
    try:
        load_price_data()
    except Exception:
        app.logger.exception("Could not warm the price data cache")


threading.Thread(target=warm_price_data, name="warm-price-data", daemon=True).start()


@app.route("/")
def index() -> str:
    filters = parse_filters()
    data = load_dashboard(
        search=filters["search"],
        selected_brands=filters["brands"],
        page=filters["page"],
        per_page=filters["per_page"],
        status=filters["status"],
        sort=filters["sort"],
        sort_dir=filters["sort_dir"],
    )
    return render_template(
        "dashboard.html",
        search=filters["search"],
        selected_per_page=data["pagination"]["per_page"],
        active_filters=build_active_filters(filters),
        base_query=build_query(filters),
        **data,
    )


def export_rows(scope: str, model: str = "") -> tuple[list[dict[str, Any]], str]:
    """Rows for an export plus a short name for the file.

    scope "all"   -> every tracked price
    scope "view"  -> whatever the dashboard filters currently show
    scope "model" -> one phone, from every source (official + Cashify)
    """
    filters = parse_filters()
    if scope == "view":
        data = load_dashboard(
            search=filters["search"], selected_brands=filters["brands"], status=filters["status"],
            sort="model", sort_dir="asc", paginate=False,
        )
        return [item[0] for item in data["latest"]], "current_view"
    data = load_dashboard(sort="model", sort_dir="asc", paginate=False)
    rows = [item[0] for item in data["latest"]]
    if scope != "model" or not model.strip():
        return rows, "all_models"
    wanted = model_key(model)
    wanted_loose = model_key(model, loose=True)
    chosen = [row for row in rows if row["model_key"] == wanted]
    if not chosen:
        chosen = [row for row in rows if model_key(row["model"], loose=True) == wanted_loose]
    if not chosen:
        chosen = [row for row in rows if matches_search(row, model)]
    label = re.sub(r"[^a-z0-9]+", "_", clean_label(model).lower()).strip("_") or "model"
    return chosen, label


@app.get("/models.json")
def model_options_json() -> Any:
    """Model names for the export picker, fetched when the menu first opens so
    the list is not embedded in every page."""
    return jsonify(load_price_data()["model_options"])


@app.get("/export.csv")
def export_csv_now() -> Any:
    rows, label = export_rows(request.args.get("scope", "view"), request.args.get("model", ""))
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["brand", "model", "part", "source", "price", "currency", "last_checked", "url"])
    for row in rows:
        writer.writerow([
            row["brand"], clean_label(row["model"]), row["part"], row["source"], row["price_value"],
            row["currency"], row["date"], row["url"],
        ])
    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=spareprice_{label}.csv"},
    )


def excel_datetime(value: str | None) -> Any:
    """ISO timestamp -> naive local datetime, which Excel can sort and filter."""
    try:
        parsed = datetime.fromisoformat((value or "").replace("Z", "+00:00"))
    except ValueError:
        return value or ""
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(DISPLAY_TIMEZONE).replace(tzinfo=None)


def style_sheet(sheet: Any, widths: list[int], money_columns: list[int], date_columns: list[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F3A8A")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for column in money_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = "#,##0"
    for column in date_columns:
        for cell in sheet[get_column_letter(column)][1:]:
            cell.number_format = "dd-mmm-yyyy hh:mm"
    sheet.freeze_panes = "A2"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions


@app.get("/export.xlsx")
def export_excel_now() -> Any:
    from openpyxl import Workbook

    rows, label = export_rows(request.args.get("scope", "all"), request.args.get("model", ""))
    workbook = Workbook()

    prices = workbook.active
    prices.title = "Prices"
    prices.append([
        "Brand", "Model", "Spare Part", "Category", "Source", "Price", "Currency",
        "Other Source", "Other Source Part", "Other Source Price", "Last Checked", "Source URL",
    ])
    for row in rows:
        other = row["competitors"][0] if row["competitors"] else None
        prices.append([
            row["brand"], clean_label(row["model"]), row["part"], row["category"] or "", row["source"],
            row["price_value"], row["currency"] or "",
            other["source"] if other else "", other["part"] if other else "",
            other["price_value"] if other else None,
            excel_datetime(row["date"]), row["url"],
        ])
    style_sheet(prices, [12, 30, 34, 16, 12, 12, 10, 14, 30, 18, 20, 60], money_columns=[6, 10], date_columns=[11])

    # One line per repair that both the brand and Cashify price, side by side.
    compare = workbook.create_sheet("Official vs Cashify")
    compare.append([
        "Brand", "Model", "Category", "Official Part", "Official Price",
        "Cashify Part", "Cashify Price", "Difference (Official - Cashify)", "Cheaper",
    ])
    for row in rows:
        if row["source"] == "Cashify":
            continue
        for other in row["competitors"]:
            if other["source"] != "Cashify" or row["price_value"] is None or other["price_value"] is None:
                continue
            difference = row["price_value"] - other["price_value"]
            cheaper = "Same" if difference == 0 else ("Cashify" if difference > 0 else row["source"])
            compare.append([
                row["brand"], clean_label(row["model"]), row["category"], row["part"], row["price_value"],
                other["part"], other["price_value"], difference, cheaper,
            ])
    style_sheet(compare, [12, 30, 16, 34, 16, 28, 16, 30, 12], money_columns=[5, 7, 8], date_columns=[])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=spareprice_{label}.xlsx"},
    )


def parse_positive_int(value: str | None, default: int) -> int:
    try:
        parsed = int(value or "")
    except ValueError:
        return default
    return parsed if parsed > 0 else default


@app.post("/check-now")
def check_now() -> Any:
    if HOSTED_READ_ONLY:
        return hosted_read_only_response()
    started = start_job("shortlist")
    if request.headers.get("Accept") == "application/json":
        return jsonify(current_job_state()), (202 if started else 409)
    return redirect(url_for("index"))


@app.post("/discover-all")
def discover_all_now() -> Any:
    if HOSTED_READ_ONLY:
        return hosted_read_only_response()
    started = start_job("catalog", brand="all", model="")
    if request.headers.get("Accept") == "application/json":
        return jsonify(current_job_state()), (202 if started else 409)
    return redirect(url_for("index"))


@app.post("/discover-scope")
def discover_scope_now() -> Any:
    if HOSTED_READ_ONLY:
        return hosted_read_only_response()
    brand = request.form.get("scrape_brand", "").strip().lower() or "all"
    model = request.form.get("scrape_model", "").strip()
    if brand not in SCRAPE_BRANDS:
        brand = "all"
    started = start_job("catalog", brand=brand, model=model)
    if request.headers.get("Accept") == "application/json":
        return jsonify(current_job_state()), (202 if started else 409)
    return redirect(url_for("index"))


@app.get("/job-status")
def job_status() -> Any:
    return jsonify(current_job_state())


def hosted_read_only_response() -> Any:
    payload = current_job_state()
    payload["message"] = "Hosted dashboard is read-only"
    payload["output"] = "Run the scraper locally, commit the updated database, then push to refresh Vercel."
    if request.headers.get("Accept") == "application/json":
        return jsonify(payload), 409
    return redirect(url_for("index"))


def update_job_progress(line: str, output: str) -> None:
    progress_patch: dict[str, Any] = {}
    discovered = re.search(
        r"Discovered\s+(?P<brand>Apple|Asus|Google|Honor|Infinix|Nothing|Nokia|OPPO|POCO|Samsung|Xiaomi|realme|OnePlus|Mi|vivo|iQOO|Motorola|Cashify)\s+(?P<model>.+?)\s+\((?:(?P<rows>\d+)\s+rows,\s+)?(?P<count>\d+)/(?P<total>[^)]+)\)",
        line,
    )
    if discovered:
        progress_patch["done"] = int(discovered.group("count"))
        brand_name = discovered.group("brand")
        model_name = discovered.group("model")
        progress_patch["current"] = model_name if model_name.lower().startswith(brand_name.lower()) else f"{brand_name} {model_name}"
        if discovered.group("rows"):
            progress_patch["rows_delta"] = int(discovered.group("rows"))

    checking = re.search(r"Checking\s+(?P<count>\d+)/(?P<total>\d+):\s+(?P<current>.+)$", line)
    if checking:
        progress_patch["done"] = int(checking.group("count")) - 1
        progress_patch["current"] = checking.group("current")

    saved = re.search(r"Saved\s+(?P<brand>\S+)\s+(?P<model>.+?):", line)
    if saved:
        progress_patch["done_increment"] = 1
        progress_patch["rows_delta"] = 1
        progress_patch["current"] = f"{saved.group('brand')} {saved.group('model')}"

    if " ERROR " in f" {line} " or " discovery failed" in line.lower() or " failed:" in line.lower():
        progress_patch["errors_delta"] = 1

    with job_lock:
        progress = dict(job_state.get("progress") or {})
        if "done" in progress_patch:
            progress["done"] = progress_patch["done"]
        if "done_increment" in progress_patch:
            progress["done"] = int(progress.get("done") or 0) + int(progress_patch["done_increment"])
        if "rows_delta" in progress_patch:
            progress["rows"] = int(progress.get("rows") or 0) + int(progress_patch["rows_delta"])
        if "errors_delta" in progress_patch:
            progress["errors"] = int(progress.get("errors") or 0) + int(progress_patch["errors_delta"])
        if "current" in progress_patch:
            progress["current"] = progress_patch["current"]
        job_state["progress"] = progress
        job_state["output"] = output


def current_job_state() -> dict[str, Any]:
    with job_lock:
        return dict(job_state)


def start_job(job_type: str, brand: str = "all", model: str = "") -> bool:
    with job_lock:
        if job_state["running"]:
            return False
        message = "Checking all configured devices..."
        if job_type == "catalog":
            scope = "all brands" if brand == "all" else brand
            if model:
                scope += f" matching {model}"
            message = f"Discovering {scope} model and spare-part prices..."
        job_state.update(
            {
                "id": str(uuid.uuid4()),
                "running": True,
                "started_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                "finished_at": None,
                "returncode": None,
                "message": message,
                "output": "",
                "progress": {
                    "done": 0,
                    "rows": 0,
                    "errors": 0,
                    "current": "",
                    "scope": message,
                },
            }
        )
    thread = threading.Thread(target=run_tracker_command, args=(job_type, brand, model), daemon=True)
    thread.start()
    return True


def run_tracker_command(job_type: str, brand: str = "all", model: str = "") -> None:
    if job_type == "catalog":
        command = [sys.executable, str(ROOT / "spareprice.py"), "discover-all", "--brand", brand, "--delay", "1"]
        if model:
            command.extend(["--model", model])
    else:
        command = [sys.executable, str(ROOT / "spareprice.py"), "run", "--config", str(DEFAULT_CONFIG)]
    try:
        process = subprocess.Popen(
            command,
            cwd=ROOT,
            text=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            bufsize=1,
        )
        output_lines = []
        assert process.stdout is not None
        for line in process.stdout:
            clean = line.strip()
            if not clean:
                continue
            output_lines.append(clean)
            if len(output_lines) > 200:
                output_lines = output_lines[-200:]
            update_job_progress(clean, "\n".join(output_lines)[-6000:])

        returncode = process.wait(timeout=3600)
        output = "\n".join(output_lines).strip()
        message = "Price check complete" if returncode == 0 else "Price check finished with errors"
        if job_type == "catalog":
            message = "Catalog discovery complete" if returncode == 0 else "Catalog discovery finished with errors"
        with job_lock:
            job_state.update(
                {
                    "running": False,
                    "finished_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "returncode": returncode,
                    "message": message,
                    "output": output[-6000:],
                }
            )
    except Exception as exc:
        with job_lock:
            job_state.update(
                {
                    "running": False,
                    "finished_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "returncode": -1,
                    "message": f"Price check failed: {type(exc).__name__}",
                    "output": str(exc),
                }
            )


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True, use_reloader=False)
